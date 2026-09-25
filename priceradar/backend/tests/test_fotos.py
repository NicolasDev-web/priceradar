"""Testes da extração de fotos e de como elas atravessam o pipeline.

Sem captura real dos portais ainda (a F1.1 do PLANO_MELHORIAS.md fica para
quando houver acesso): os casos cobrem os formatos do schema.org e o molde de
URL do CDN do Grupo ZAP, que é o que os scrapers vão entregar a `extrair_fotos`.
"""
import sys
from io import BytesIO
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from models import Empreendimento  # noqa: E402
from scraper.parser import MAX_FOTOS, extrair_fotos, fotos_de_card  # noqa: E402
from services.deduplicador import deduplicar_cross_portal  # noqa: E402
from services.export import COL, gerar_excel  # noqa: E402


def test_string_unica():
    assert extrair_fotos("https://img.exemplo.com/a.jpg") == ["https://img.exemplo.com/a.jpg"]


def test_lista_e_image_object():
    valor = [
        "https://img.exemplo.com/a.jpg",
        {"@type": "ImageObject", "url": "https://img.exemplo.com/b.jpg"},
        {"@type": "ImageObject", "contentUrl": "https://img.exemplo.com/c.jpg"},
    ]
    assert extrair_fotos(valor) == [
        "https://img.exemplo.com/a.jpg",
        "https://img.exemplo.com/b.jpg",
        "https://img.exemplo.com/c.jpg",
    ]


def test_molde_do_cdn_grupo_zap_e_preenchido():
    molde = "https://resizedimgs.vivareal.com/{action}/{width}x{height}/named.images.sp/abc/foto.jpg"
    [url] = extrair_fotos(molde)
    assert "{" not in url
    assert url == "https://resizedimgs.vivareal.com/fit-in/870x653/named.images.sp/abc/foto.jpg"


def test_relativa_e_sem_protocolo():
    assert extrair_fotos("/fotos/1.jpg", "https://www.portal.com.br") == ["https://www.portal.com.br/fotos/1.jpg"]
    assert extrair_fotos("//cdn.portal.com/1.jpg") == ["https://cdn.portal.com/1.jpg"]
    # Relativa sem base não tem como virar URL — descarta em vez de inventar.
    assert extrair_fotos("/fotos/1.jpg") == []


def test_descarta_lixo_duplicata_e_respeita_teto():
    valor = [
        "data:image/gif;base64,R0lGOD",
        "https://img.exemplo.com/logo-imobiliaria.png",
        "https://img.exemplo.com/icone.svg",
        "https://img.exemplo.com/a.jpg",
        "https://img.exemplo.com/a.jpg",
    ] + [f"https://img.exemplo.com/{i}.jpg" for i in range(20)]
    fotos = extrair_fotos(valor)
    assert fotos[0] == "https://img.exemplo.com/a.jpg"
    assert len(fotos) == MAX_FOTOS
    assert len(set(fotos)) == len(fotos)


def test_formato_inesperado_nao_explode():
    assert extrair_fotos(None) == []
    assert extrair_fotos(42) == []
    assert extrair_fotos({"sem": "url"}) == []


def test_card_html_prefere_lazy_load():
    html = """
    <div class="card">
      <img src="data:image/gif;base64,xx" data-src="https://img.exemplo.com/1.jpg">
      <img src="https://img.exemplo.com/2.jpg">
      <picture><source srcset="https://img.exemplo.com/p.jpg 320w, https://img.exemplo.com/g.jpg 1024w"></picture>
      <img src="/static/logo.svg">
    </div>"""
    card = BeautifulSoup(html, "lxml").select_one(".card")
    assert fotos_de_card(card, "https://www.portal.com.br") == [
        "https://img.exemplo.com/1.jpg",
        "https://img.exemplo.com/2.jpg",
        "https://img.exemplo.com/g.jpg",
    ]


def _anuncio(portal, url, fotos, **kw):
    base = {
        "id": url, "nome_anuncio": "Apartamento 2 quartos Meireles", "nome_empreendimento": None,
        "construtora": None, "cidade": "fortaleza", "bairro": "Meireles", "portal": portal,
        "preco": 400_000.0, "area_m2": 60.0, "preco_m2": 6_666.67, "quartos": 2,
        "banheiros": 2, "vagas": 1, "descricao": None, "url_anuncio": url,
        "data_coleta": datetime.now(), "fotos": fotos,
    }
    base.update(kw)
    return base


def test_dedup_une_fotos_dos_portais():
    a = _anuncio("vivareal", "https://vr/1", ["https://img/a.jpg", "https://img/b.jpg"],
                 latitude=-3.73, longitude=-38.5)
    b = _anuncio("zapimoveis", "https://zap/1", ["https://img/b.jpg", "https://img/c.jpg"],
                 latitude=-3.73, longitude=-38.5)
    [unico] = deduplicar_cross_portal([a, b])
    assert sorted(unico["fotos"]) == ["https://img/a.jpg", "https://img/b.jpg", "https://img/c.jpg"]


def test_export_traz_link_da_foto(tmp_path):
    emp = Empreendimento(**_anuncio("vivareal", "https://vr/1", ["https://img/capa.jpg", "https://img/2.jpg"]))
    sem = Empreendimento(**_anuncio("vivareal", "https://vr/2", []))
    ws = load_workbook(BytesIO(gerar_excel([emp, sem], 6_666.67)))["ANUNCIOS"]
    celulas = {c.value: c for c in ws[COL["Foto"]][1:]}
    assert celulas["Ver foto"].hyperlink.target == "https://img/capa.jpg"
    assert None in celulas or "" in celulas
